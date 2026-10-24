from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile

from flask import Blueprint, abort, current_app, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from sqlalchemy import case, or_
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import joinedload

from app.extensions import db
from app.services.agenda_excel import AgendaExcelError, import_agenda_workbook
from app.models import (
    AgendamentoLaboratorio,
    HistoricoAgendamentoLaboratorio,
    Laboratorio,
    NotificacaoAgenda,
    MensagemAgendamento,
    NotificacaoSistema,
    Recado,
    Usuario,
)
from app.utils.authz import (
    SUPPORT_ROLES,
    has_role,
    is_master_admin,
    master_admin_required,
    roles_required,
)
from app.services.notifications import notify_support as notify_general_support
from app.utils.validation import (
    ValidationError,
    choice,
    integer,
    optional_id,
    optional_text,
    required_text,
)
from app.utils.export import excel_safe


agenda_bp = Blueprint("agenda", __name__, url_prefix="/agendamentos")
PRIORIDADES_RECADOS = {"normal", "importante", "urgente"}
TURNOS = {
    "manha": {"label": "Manhã", "inicio": time(8, 0), "fim": time(12, 0)},
    "tarde": {"label": "Tarde", "inicio": time(13, 0), "fim": time(17, 0)},
    "noite": {"label": "Noite", "inicio": time(18, 0), "fim": time(22, 0)},
}
STATUS_ATIVOS = {"solicitado", "agendado"}
STATUS_LABELS = {
    "solicitado": "Em análise",
    "agendado": "Confirmado",
    "recusado": "Recusado",
    "cancelado": "Cancelado",
    "excluido": "Excluído",
}
ILHAS_LABORATORIO = {
    "ilha-a": {"nome": "Ilha A", "laboratorios": ("Lab 05", "Lab 06"), "capacidade": 56},
    "ilha-b": {"nome": "Ilha B", "laboratorios": ("Lab 07", "Lab 08"), "capacidade": 56},
}


def utc_now():
    return datetime.now(timezone.utc).replace(tzinfo=None)


def report_period():
    """Retorna o intervalo solicitado para os relatórios de agendamento."""
    periodo = (request.args.get("periodo") or "dia").strip().lower()
    if periodo not in {"dia", "mes", "ano"}:
        raise ValidationError("Período de relatório inválido.")
    referencia = parse_date_value(request.args.get("data") or date.today().isoformat(), "Data")
    if periodo == "dia":
        inicio = referencia
        fim = referencia + timedelta(days=1)
        rotulo = referencia.strftime("%d-%m-%Y")
        titulo = f"Agendamentos de {referencia.strftime('%d/%m/%Y')}"
    elif periodo == "mes":
        inicio = referencia.replace(day=1)
        fim = (inicio.replace(day=28) + timedelta(days=4)).replace(day=1)
        rotulo = inicio.strftime("%m-%Y")
        titulo = f"Agendamentos de {inicio.strftime('%m/%Y')}"
    else:
        inicio = referencia.replace(month=1, day=1)
        fim = inicio.replace(year=inicio.year + 1)
        rotulo = str(inicio.year)
        titulo = f"Agendamentos de {inicio.year}"
    return periodo, inicio, fim, rotulo, titulo


def report_bookings(inicio, fim):
    return (
        AgendamentoLaboratorio.query.options(
            joinedload(AgendamentoLaboratorio.laboratorio),
            joinedload(AgendamentoLaboratorio.laboratorio_secundario),
            joinedload(AgendamentoLaboratorio.usuario),
            joinedload(AgendamentoLaboratorio.atendido_por),
            joinedload(AgendamentoLaboratorio.confirmado_por),
        )
        .filter(
            AgendamentoLaboratorio.status != "excluido",
            AgendamentoLaboratorio.inicio >= datetime.combine(inicio, time.min),
            AgendamentoLaboratorio.inicio < datetime.combine(fim, time.min),
        )
        .order_by(AgendamentoLaboratorio.inicio, AgendamentoLaboratorio.laboratorio_id)
        .all()
    )


def parse_date_value(value, label):
    try:
        return date.fromisoformat((value or "").strip())
    except (TypeError, ValueError):
        raise ValidationError(f"{label} inválida.")


def selected_date():
    value = (request.args.get("data") or "").strip()
    if not value:
        return date.today()
    try:
        return date.fromisoformat(value)
    except ValueError:
        return date.today()


def can_manage_booking(booking):
    return has_role(*SUPPORT_ROLES) or booking.usuario_id == current_user.id


def booking_uses_laboratorio(booking, laboratorio_id):
    return laboratorio_id in {booking.laboratorio_id, booking.laboratorio_secundario_id}


def get_ilha(codigo, exigir_disponivel=True):
    configuracao = ILHAS_LABORATORIO.get(codigo)
    if not configuracao:
        raise ValidationError("Ilha de laboratórios inválida.")
    laboratorios = [
        Laboratorio.query.filter_by(nome=nome, ativo=True).first()
        for nome in configuracao["laboratorios"]
    ]
    if exigir_disponivel and not all(laboratorios):
        raise ValidationError("Esta ilha não está disponível para agendamento.")
    return configuracao, laboratorios


def register_history(booking, action, description=None, actor=None):
    db.session.add(
        HistoricoAgendamentoLaboratorio(
            agendamento=booking,
            usuario_id=(actor or current_user).id,
            acao=action,
            descricao=description,
        )
    )


def claim_booking(booking, allow_master_override=False):
    """Atribui a solicitação ao técnico atual ou valida o responsável já definido."""
    if not current_user.matricula:
        raise ValidationError(
            "Cadastre sua matrícula antes de tratar solicitações de agendamento."
        )
    if booking.atendido_por_id and booking.atendido_por_id != current_user.id:
        if allow_master_override and is_master_admin():
            return
        responsavel = booking.atendido_por
        raise ValidationError(
            f"Esta solicitação já está em atendimento por {responsavel.nome} ({responsavel.matricula})."
        )
    if not booking.atendido_por_id:
        booking.atendido_por_id = current_user.id
        booking.atendido_em = utc_now()
        register_history(
            booking,
            "atendimento_assumido",
            f"Atendimento assumido por {current_user.nome} ({current_user.matricula}).",
        )


def notify_support(booking):
    possui_software = bool(booking.softwares_necessarios)
    complemento_software = " Há solicitação de software." if possui_software else ""
    recipients = Usuario.query.filter(
        Usuario.ativo.is_(True), Usuario.perfil.in_(SUPPORT_ROLES)
    ).all()
    for recipient in recipients:
        db.session.add(
            NotificacaoAgenda(
                usuario_id=recipient.id,
                agendamento=booking,
                mensagem=(
                    f"Nova solicitação {booking.protocolo}: {booking.professor_nome} "
                    f"solicitou {booking.laboratorio_label}.{complemento_software}"
                ),
                tipo="software" if possui_software else "solicitacao",
            )
        )
    notify_general_support(
        "agendamento",
        "Nova solicitação de agendamento",
        f"{booking.professor_nome} solicitou {booking.laboratorio_label} ({booking.protocolo}).{complemento_software}",
        url_for("agenda.index", data=booking.inicio.date().isoformat()),
    )


def resolve_booking_notifications(booking):
    """Marca como lidos os avisos de uma solicitação que já recebeu tratativa."""
    resolved_at = utc_now()
    NotificacaoAgenda.query.filter_by(agendamento_id=booking.id, lida_em=None).update(
        {NotificacaoAgenda.lida_em: resolved_at}, synchronize_session=False
    )
    NotificacaoSistema.query.filter(
        NotificacaoSistema.tipo == "agendamento",
        NotificacaoSistema.lida_em.is_(None),
        NotificacaoSistema.mensagem.contains(booking.protocolo),
    ).update({NotificacaoSistema.lida_em: resolved_at}, synchronize_session=False)


def booking_conflict(laboratorio_id, inicio, fim, ignore_id=None):
    query = AgendamentoLaboratorio.query.filter(
        or_(
            AgendamentoLaboratorio.laboratorio_id == laboratorio_id,
            AgendamentoLaboratorio.laboratorio_secundario_id == laboratorio_id,
        ),
        AgendamentoLaboratorio.status == "agendado",
        AgendamentoLaboratorio.inicio < fim,
        AgendamentoLaboratorio.fim > inicio,
    )
    if ignore_id is not None:
        query = query.filter(AgendamentoLaboratorio.id != ignore_id)
    return query.with_for_update().first()


def apply_booking_fields(booking):
    data_reserva = parse_date_value(request.form.get("data_reserva"), "Data da reserva")
    if data_reserva.weekday() == 6:
        raise ValidationError("Não é possível agendar laboratórios aos domingos.")
    turno_codigo = choice(request.form, "turno", set(TURNOS), "Turno")
    if data_reserva.weekday() == 5 and turno_codigo == "noite":
        raise ValidationError("Não é possível agendar laboratórios no turno da noite aos sábados.")
    turno = TURNOS[turno_codigo]
    inicio = datetime.combine(data_reserva, turno["inicio"])
    fim = datetime.combine(data_reserva, turno["fim"])

    reserva_interna = has_role(*SUPPORT_ROLES)
    quantidade_informada = (request.form.get("quantidade_alunos") or "").strip()
    quantidade_alunos = 0 if reserva_interna and not quantidade_informada else integer(
        request.form, "quantidade_alunos", "Quantidade de alunos", 1
    )
    ilha_codigo = (request.form.get("ilha_codigo") or "").strip()
    if quantidade_alunos > 56:
        raise ValidationError("A capacidade máxima por ilha é de 56 alunos.")

    if ilha_codigo:
        if quantidade_alunos < 33:
            raise ValidationError("A reserva de uma ilha é destinada a turmas de 33 a 56 alunos.")
        configuracao_ilha, laboratorios_reservados = get_ilha(ilha_codigo)
    else:
        if quantidade_alunos > 32:
            raise ValidationError(
                "Para turmas de 33 a 56 alunos, escolha uma ilha de laboratórios disponível."
            )
        laboratorio_id = optional_id(request.form, "laboratorio_id")
        laboratorio = Laboratorio.query.filter_by(id=laboratorio_id, ativo=True).first()
        if not laboratorio:
            raise ValidationError("Laboratório inválido ou indisponível.")
        configuracao_ilha = None
        laboratorios_reservados = [laboratorio]

    capacidade_total = sum(laboratorio.capacidade or 0 for laboratorio in laboratorios_reservados)
    if capacidade_total and quantidade_alunos > capacidade_total:
        raise ValidationError(
            f"A quantidade de alunos excede a capacidade de {capacidade_total} computadores desta opção."
        )

    if configuracao_ilha:
        for laboratorio in laboratorios_reservados:
            if booking.status == "agendado" and booking_conflict(laboratorio.id, inicio, fim, booking.id):
                raise ValidationError(
                    f"{laboratorio.nome} já está reservado no turno selecionado. Escolha outra opção."
                )
    else:
        for laboratorio in laboratorios_reservados:
            if booking.status == "agendado" and booking_conflict(laboratorio.id, inicio, fim, booking.id):
                raise ValidationError(
                    f"{laboratorio.nome} já está reservado no turno selecionado. Escolha outra opção."
                )

    booking.laboratorio_id = laboratorios_reservados[0].id
    booking.laboratorio_secundario_id = (
        laboratorios_reservados[1].id if len(laboratorios_reservados) > 1 else None
    )
    booking.ilha_codigo = ilha_codigo or None
    if reserva_interna:
        booking.professor_nome = required_text(request.form, "responsavel_nome", "Nome do responsável", 120)
        booking.professor_matricula = optional_text(request.form, "professor_matricula", 60) or ""
    else:
        booking.professor_nome = required_text(request.form, "professor_nome", "Nome do professor", 120)
        booking.professor_matricula = required_text(request.form, "professor_matricula", "Matrícula", 60)
    booking.quantidade_alunos = quantidade_alunos
    booking.titulo = required_text(request.form, "titulo", "Título", 150)
    booking.finalidade = optional_text(request.form, "finalidade")
    booking.softwares_necessarios = optional_text(request.form, "softwares_necessarios")
    booking.inicio = inicio
    booking.fim = fim


@agenda_bp.route("/")
@login_required
def index():
    if current_user.perfil == "professor":
        return redirect(url_for("agendamento_laboratorios"))
    data_selecionada = selected_date()
    inicio_dia = datetime.combine(data_selecionada, time.min)
    fim_dia = inicio_dia + timedelta(days=1)

    agendamentos = (
        AgendamentoLaboratorio.query.options(
            joinedload(AgendamentoLaboratorio.laboratorio),
            joinedload(AgendamentoLaboratorio.laboratorio_secundario),
            joinedload(AgendamentoLaboratorio.usuario),
        )
        .filter(
            AgendamentoLaboratorio.status.in_(STATUS_ATIVOS),
            AgendamentoLaboratorio.inicio < fim_dia,
            AgendamentoLaboratorio.fim > inicio_dia,
        )
        .order_by(AgendamentoLaboratorio.inicio, AgendamentoLaboratorio.laboratorio_id)
        .all()
    )
    recados = (
        Recado.query.options(joinedload(Recado.autor))
        .filter(Recado.data_inicio <= data_selecionada, Recado.data_fim >= data_selecionada)
        .order_by(
            case(
                (Recado.prioridade == "urgente", 0),
                (Recado.prioridade == "importante", 1),
                else_=2,
            ),
            Recado.criado_em.desc(),
        )
        .all()
    )
    laboratorios = Laboratorio.query.filter_by(ativo=True).order_by(Laboratorio.nome).all()
    mapa_turnos = []
    disponibilidade_interna = {}
    for codigo, turno in TURNOS.items():
        inicio_turno = datetime.combine(data_selecionada, turno["inicio"])
        fim_turno = datetime.combine(data_selecionada, turno["fim"])
        slots = []
        for laboratorio in laboratorios:
            reservas_slot = [
                item
                for item in agendamentos
                if booking_uses_laboratorio(item, laboratorio.id)
                and item.inicio < fim_turno
                and item.fim > inicio_turno
            ]
            slots.append({"laboratorio": laboratorio, "reservas": reservas_slot})
        ilhas = []
        for codigo_ilha, configuracao_ilha in ILHAS_LABORATORIO.items():
            _, laboratorios_ilha = get_ilha(codigo_ilha, exigir_disponivel=False)
            if not all(laboratorios_ilha):
                continue
            disponivel = all(
                not any(
                    booking_uses_laboratorio(item, laboratorio.id)
                    and item.inicio < fim_turno
                    and item.fim > inicio_turno
                    and item.status == "agendado"
                    for item in agendamentos
                )
                for laboratorio in laboratorios_ilha
            )
            ilhas.append({"codigo": codigo_ilha, "configuracao": configuracao_ilha, "laboratorios": laboratorios_ilha, "disponivel": disponivel})
        mapa_turnos.append(
            {
                "codigo": codigo,
                "label": turno["label"],
                "periodo": f"{turno['inicio'].strftime('%H:%M')}–{turno['fim'].strftime('%H:%M')}",
                "slots": slots,
                "ilhas": ilhas,
            }
        )
        disponibilidade_interna[codigo] = {
            "laboratorios_disponiveis": [
                slot["laboratorio"].id
                for slot in slots
                if not any(reserva.status == "agendado" for reserva in slot["reservas"])
            ],
            "ilhas_disponiveis": {
                ilha["codigo"]: ilha["disponivel"] for ilha in ilhas
            },
        }

    preselected_lab_id = request.args.get("lab", type=int)
    if not any(item.id == preselected_lab_id for item in laboratorios):
        preselected_lab_id = laboratorios[0].id if laboratorios else None
    preselected_turno = request.args.get("turno")
    if preselected_turno not in TURNOS:
        preselected_turno = "manha"
    if data_selecionada.weekday() == 5 and preselected_turno == "noite":
        preselected_turno = "manha"
    todos_laboratorios = []
    if has_role(*SUPPORT_ROLES):
        todos_laboratorios = Laboratorio.query.order_by(Laboratorio.nome).all()
    notificacoes = []
    solicitacoes_pendentes = []
    if has_role(*SUPPORT_ROLES):
        notificacoes = (
            NotificacaoAgenda.query.options(joinedload(NotificacaoAgenda.agendamento))
            .filter_by(usuario_id=current_user.id, lida_em=None)
            .order_by(NotificacaoAgenda.criado_em.desc())
            .limit(10)
            .all()
        )
        solicitacoes_pendentes = (
            AgendamentoLaboratorio.query.options(
                joinedload(AgendamentoLaboratorio.laboratorio),
                joinedload(AgendamentoLaboratorio.usuario),
            )
            .filter_by(status="solicitado")
            .order_by(AgendamentoLaboratorio.criado_em, AgendamentoLaboratorio.id)
            .all()
        )

    return render_template(
        "agenda/index.html",
        agendamentos=agendamentos,
        recados=recados,
        laboratorios=laboratorios,
        mapa_turnos=mapa_turnos,
        disponibilidade_interna=disponibilidade_interna,
        turnos=TURNOS,
        preselected_lab_id=preselected_lab_id,
        preselected_turno=preselected_turno,
        todos_laboratorios=todos_laboratorios,
        data_selecionada=data_selecionada,
        data_anterior=data_selecionada - timedelta(days=1),
        proxima_data=data_selecionada + timedelta(days=1),
        hoje=date.today(),
        domingo=data_selecionada.weekday() == 6,
        sabado=data_selecionada.weekday() == 5,
        suporte=has_role(*SUPPORT_ROLES),
        status_labels=STATUS_LABELS,
        notificacoes=notificacoes,
        solicitacoes_pendentes=solicitacoes_pendentes,
    )


@agenda_bp.route("/relatorios")
@roles_required("admin", "tecnico")
def relatorios():
    periodo = (request.args.get("periodo") or "dia").strip().lower()
    if periodo not in {"dia", "mes", "ano"}:
        periodo = "dia"
    try:
        referencia = parse_date_value(request.args.get("data") or date.today().isoformat(), "Data")
    except ValidationError:
        referencia = date.today()
    return render_template("agenda/relatorios.html", periodo=periodo, referencia=referencia)


@agenda_bp.route("/importar-excel", methods=["GET", "POST"])
@master_admin_required
def importar_excel():
    """Importa reservas futuras dos laboratórios 5 a 8 a partir de um mapa Excel."""
    result = None
    starts_on = date.today()

    if request.method == "POST":
        arquivo = request.files.get("arquivo")
        try:
            starts_on = parse_date_value(request.form.get("a_partir_de"), "Data inicial")
        except ValidationError as error:
            flash(str(error), "danger")
            return redirect(url_for("agenda.importar_excel"))

        if not arquivo or not arquivo.filename:
            flash("Selecione a planilha de agendamentos.", "danger")
            return redirect(url_for("agenda.importar_excel"))
        if Path(arquivo.filename).suffix.lower() != ".xlsx":
            flash("Envie uma planilha no formato .xlsx.", "danger")
            return redirect(url_for("agenda.importar_excel"))

        temporary_path = None
        try:
            with NamedTemporaryFile(
                suffix=".xlsx", dir=current_app.config["UPLOAD_FOLDER"], delete=False
            ) as temporary_file:
                arquivo.save(temporary_file)
                temporary_path = Path(temporary_file.name)
            confirm = request.form.get("acao") == "confirmar"
            result = import_agenda_workbook(
                temporary_path,
                user_email=current_user.email,
                lab_names={"Lab 05", "Lab 06", "Lab 07", "Lab 08"},
                starts_on=starts_on,
                commit=confirm,
            )
            if confirm:
                flash(f"{result.importados} agendamento(s) foram importados.", "success")
            else:
                flash("Simulação concluída. Nenhum agendamento foi gravado.", "info")
        except AgendaExcelError as error:
            flash(str(error), "danger")
        finally:
            if temporary_path:
                temporary_path.unlink(missing_ok=True)

    return render_template(
        "agenda/importar_excel.html",
        starts_on=starts_on,
        result=result,
    )


@agenda_bp.route("/relatorios/excel")
@roles_required("admin", "tecnico")
def relatorio_excel():
    try:
        _periodo, inicio, fim, rotulo, titulo = report_period()
    except ValidationError as error:
        flash(str(error), "danger")
        return redirect(url_for("agenda.relatorios"))

    agendamentos = report_bookings(inicio, fim)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Agendamentos"
    worksheet.append([titulo])
    worksheet.append([f"Período: {inicio.strftime('%d/%m/%Y')} a {(fim - timedelta(days=1)).strftime('%d/%m/%Y')}"])
    worksheet.append([])
    headers = [
        "Protocolo", "Data", "Turno", "Laboratório", "Professor", "Matrícula",
        "Alunos", "Atividade", "Status", "Responsável", "Matrícula responsável",
        "Solicitado em", "Confirmado em",
    ]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3B70")
    for cell in worksheet[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4169E1")
    for booking in agendamentos:
        worksheet.append(
            [
                booking.protocolo,
                booking.inicio.date(),
                booking.turno_label,
                excel_safe(booking.laboratorio_label),
                excel_safe(booking.professor_nome or booking.usuario.nome),
                excel_safe(booking.professor_matricula),
                booking.quantidade_alunos,
                excel_safe(booking.titulo),
                STATUS_LABELS.get(booking.status, booking.status),
                excel_safe(booking.atendido_por.nome) if booking.atendido_por else "",
                excel_safe(booking.atendido_por.matricula) if booking.atendido_por else "",
                booking.criado_em.strftime("%d/%m/%Y %H:%M") if booking.criado_em else "",
                booking.confirmado_em.strftime("%d/%m/%Y %H:%M") if booking.confirmado_em else "",
            ]
        )
    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = f"A4:M{max(4, worksheet.max_row)}"
    for cell in worksheet["B"][4:]:
        cell.number_format = "dd/mm/yyyy"
    widths = [16, 13, 12, 16, 28, 16, 10, 40, 16, 26, 22, 20, 20]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + column)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"agendamentos_{rotulo}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@agenda_bp.route("/relatorios/pdf")
@roles_required("admin", "tecnico")
def relatorio_pdf():
    try:
        _periodo, inicio, fim, rotulo, titulo = report_period()
    except ValidationError as error:
        flash(str(error), "danger")
        return redirect(url_for("agenda.relatorios"))

    agendamentos = report_bookings(inicio, fim)
    output = BytesIO()
    largura, altura = landscape(A4)
    pdf = canvas.Canvas(output, pagesize=landscape(A4))

    def cabecalho():
        y_pos = altura - 34
        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(28, y_pos, titulo)
        pdf.setFont("Helvetica", 9)
        pdf.drawString(28, y_pos - 15, f"Período: {inicio.strftime('%d/%m/%Y')} a {(fim - timedelta(days=1)).strftime('%d/%m/%Y')}")
        y_pos -= 34
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawString(28, y_pos, "Data / turno")
        pdf.drawString(105, y_pos, "Laboratório")
        pdf.drawString(180, y_pos, "Professor")
        pdf.drawString(350, y_pos, "Atividade")
        pdf.drawString(600, y_pos, "Alunos")
        pdf.drawString(652, y_pos, "Status")
        return y_pos - 13

    y = cabecalho()
    pdf.setFont("Helvetica", 8)
    for booking in agendamentos:
        if y < 34:
            pdf.showPage()
            y = cabecalho()
            pdf.setFont("Helvetica", 8)
        data_turno = f"{booking.inicio.strftime('%d/%m/%Y')}\n{booking.turno_label}"
        pdf.drawString(28, y, data_turno.replace("\n", " "))
        pdf.drawString(105, y, booking.laboratorio_label[:22])
        pdf.drawString(180, y, (booking.professor_nome or booking.usuario.nome)[:29])
        pdf.drawString(350, y, booking.titulo[:42])
        pdf.drawRightString(630, y, str(booking.quantidade_alunos))
        pdf.drawString(652, y, STATUS_LABELS.get(booking.status, booking.status)[:15])
        y -= 14
    if not agendamentos:
        pdf.drawString(28, y, "Nenhum agendamento encontrado para o período.")
    pdf.save()
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"agendamentos_{rotulo}.pdf", mimetype="application/pdf")


@agenda_bp.route("/solicitar")
@login_required
def solicitar_professor():
    if has_role(*SUPPORT_ROLES):
        return redirect(url_for("agenda.index"))
    data_selecionada = selected_date()
    inicio_dia = datetime.combine(data_selecionada, time.min)
    fim_dia = inicio_dia + timedelta(days=1)
    laboratorios = Laboratorio.query.filter_by(ativo=True).order_by(Laboratorio.nome).all()
    confirmados = (
        AgendamentoLaboratorio.query.filter(
            AgendamentoLaboratorio.status == "agendado",
            AgendamentoLaboratorio.inicio < fim_dia,
            AgendamentoLaboratorio.fim > inicio_dia,
        )
        .all()
    )
    minhas_reservas_ativas = (
        AgendamentoLaboratorio.query.filter(
            AgendamentoLaboratorio.usuario_id == current_user.id,
            AgendamentoLaboratorio.status.in_(STATUS_ATIVOS),
            AgendamentoLaboratorio.inicio < fim_dia,
            AgendamentoLaboratorio.fim > inicio_dia,
        )
        .all()
    )
    mapa_turnos = []
    disponibilidade_formulario = {}
    for codigo, turno in TURNOS.items():
        inicio_turno = datetime.combine(data_selecionada, turno["inicio"])
        fim_turno = datetime.combine(data_selecionada, turno["fim"])
        slots = []
        for laboratorio in laboratorios:
            reserva_propria = next(
                (
                    item
                    for item in minhas_reservas_ativas
                    if booking_uses_laboratorio(item, laboratorio.id)
                    and item.inicio < fim_turno
                    and item.fim > inicio_turno
                    and item.usuario_id == current_user.id
                ),
                None,
            )
            ocupado = any(
                booking_uses_laboratorio(item, laboratorio.id)
                and item.inicio < fim_turno
                and item.fim > inicio_turno
                for item in confirmados
            )
            slots.append({"laboratorio": laboratorio, "ocupado": ocupado, "reserva_propria": reserva_propria})
        ilhas = []
        for codigo_ilha, configuracao_ilha in ILHAS_LABORATORIO.items():
            _, laboratorios_ilha = get_ilha(codigo_ilha, exigir_disponivel=False)
            if not all(laboratorios_ilha):
                continue
            disponivel = all(
                not any(
                    booking_uses_laboratorio(item, laboratorio.id)
                    and item.inicio < fim_turno
                    and item.fim > inicio_turno
                    for item in confirmados
                )
                for laboratorio in laboratorios_ilha
            )
            ilhas.append({"codigo": codigo_ilha, "configuracao": configuracao_ilha, "laboratorios": laboratorios_ilha, "disponivel": disponivel})
        disponibilidade_formulario[codigo] = {
            "laboratorios_disponiveis": [
                slot["laboratorio"].id for slot in slots if not slot["ocupado"]
            ],
            "ilhas_disponiveis": {
                ilha["codigo"]: ilha["disponivel"] for ilha in ilhas
            },
        }
        mapa_turnos.append({"codigo": codigo, "label": turno["label"], "periodo": f"{turno['inicio'].strftime('%H:%M')}–{turno['fim'].strftime('%H:%M')}", "slots": slots, "ilhas": ilhas})

    preselected_lab_id = request.args.get("lab", type=int)
    if not any(laboratorio.id == preselected_lab_id for laboratorio in laboratorios):
        preselected_lab_id = laboratorios[0].id if laboratorios else None
    preselected_turno = request.args.get("turno")
    if preselected_turno not in TURNOS:
        preselected_turno = "manha"
    if data_selecionada.weekday() == 5 and preselected_turno == "noite":
        preselected_turno = "manha"
    return render_template(
        "agenda/solicitar.html",
        laboratorios=laboratorios,
        turnos=TURNOS,
        hoje=date.today(),
        professor=current_user,
        data_selecionada=data_selecionada,
        data_anterior=data_selecionada - timedelta(days=1),
        proxima_data=data_selecionada + timedelta(days=1),
        mapa_turnos=mapa_turnos,
        disponibilidade_formulario=disponibilidade_formulario,
        preselected_lab_id=preselected_lab_id,
        preselected_turno=preselected_turno,
        domingo=data_selecionada.weekday() == 6,
        sabado=data_selecionada.weekday() == 5,
    )


@agenda_bp.route("/minhas-solicitacoes")
@login_required
def minhas_solicitacoes():
    if current_user.perfil != "professor":
        abort(403)
    agendamentos = AgendamentoLaboratorio.query.options(joinedload(AgendamentoLaboratorio.laboratorio), joinedload(AgendamentoLaboratorio.laboratorio_secundario), joinedload(AgendamentoLaboratorio.atendido_por)).filter_by(usuario_id=current_user.id).order_by(AgendamentoLaboratorio.inicio.desc()).all()
    return render_template("agenda/minhas_solicitacoes.html", agendamentos=agendamentos, status_labels=STATUS_LABELS)


@agenda_bp.route("/<int:id>/detalhes")
@login_required
def detalhes_agendamento(id):
    booking = AgendamentoLaboratorio.query.options(joinedload(AgendamentoLaboratorio.laboratorio), joinedload(AgendamentoLaboratorio.laboratorio_secundario), joinedload(AgendamentoLaboratorio.atendido_por), joinedload(AgendamentoLaboratorio.historicos).joinedload(HistoricoAgendamentoLaboratorio.usuario)).get_or_404(id)
    if current_user.perfil == "professor" and booking.usuario_id != current_user.id:
        abort(403)
    if current_user.perfil not in SUPPORT_ROLES and booking.usuario_id != current_user.id:
        abort(403)
    return render_template("agenda/detalhes.html", booking=booking, status_labels=STATUS_LABELS)


@agenda_bp.route("/<int:id>/mensagens", methods=["GET", "POST"])
@login_required
def mensagens_agendamento(id):
    booking = AgendamentoLaboratorio.query.get_or_404(id)
    if current_user.perfil == "professor" and booking.usuario_id != current_user.id:
        abort(403)
    if current_user.perfil not in SUPPORT_ROLES and booking.usuario_id != current_user.id:
        abort(403)
    if request.method == "POST":
        try:
            db.session.add(MensagemAgendamento(agendamento_id=booking.id, autor_id=current_user.id, mensagem=required_text(request.form, "mensagem", "Mensagem")))
            db.session.commit()
            flash("Mensagem enviada.", "success")
        except (ValidationError, SQLAlchemyError) as error:
            db.session.rollback(); flash(str(error) if isinstance(error, ValidationError) else "Não foi possível enviar a mensagem.", "danger")
    mensagens = MensagemAgendamento.query.options(joinedload(MensagemAgendamento.autor)).filter_by(agendamento_id=booking.id).order_by(MensagemAgendamento.criado_em).all()
    return render_template("agenda/mensagens.html", booking=booking, mensagens=mensagens)


@agenda_bp.route("/novo", methods=["POST"])
@login_required
def novo_agendamento():
    reserva_interna = has_role(*SUPPORT_ROLES)
    booking = AgendamentoLaboratorio(
        usuario_id=current_user.id,
        status="agendado" if reserva_interna else "solicitado",
    )
    try:
        apply_booking_fields(booking)
        db.session.add(booking)
        db.session.flush()
        if reserva_interna:
            booking.atendido_por_id = current_user.id
            booking.atendido_em = utc_now()
            booking.confirmado_por_id = current_user.id
            booking.confirmado_em = utc_now()
            register_history(
                booking,
                "agendamento_interno_criado",
                "Agendamento interno confirmado pela equipe responsável.",
            )
        else:
            register_history(booking, "solicitacao_criada", "Solicitação aguardando análise da equipe.")
            notify_support(booking)
        db.session.commit()
    except ValidationError as error:
        db.session.rollback()
        flash(str(error), "danger")
    except SQLAlchemyError:
        db.session.rollback()
        flash("Não foi possível salvar a reserva.", "danger")
    else:
        if reserva_interna:
            flash(f"Agendamento interno {booking.protocolo} registrado e confirmado.", "success")
        else:
            flash(
                f"Solicitação {booking.protocolo} enviada. A equipe confirmará o agendamento após análise.",
                "success",
            )
    destino = "agendamento_laboratorios" if current_user.perfil == "professor" else "agenda.index"
    return redirect(url_for(destino, data=request.form.get("data_retorno") or date.today()))


@agenda_bp.route("/<int:id>/editar", methods=["GET", "POST"])
@login_required
def editar_agendamento(id):
    booking = AgendamentoLaboratorio.query.get_or_404(id)
    if not can_manage_booking(booking) or booking.status != "agendado":
        abort(403)

    laboratorios = Laboratorio.query.filter(
        (Laboratorio.ativo.is_(True)) | (Laboratorio.id == booking.laboratorio_id)
    ).order_by(Laboratorio.nome).all()
    if request.method == "POST":
        try:
            apply_booking_fields(booking)
            register_history(booking, "solicitacao_atualizada", "Dados da solicitação foram atualizados.")
            db.session.commit()
        except ValidationError as error:
            db.session.rollback()
            flash(str(error), "danger")
        except SQLAlchemyError:
            db.session.rollback()
            flash("Não foi possível atualizar a reserva.", "danger")
        else:
            flash("Reserva atualizada com sucesso.", "success")
            return redirect(url_for("agenda.index", data=booking.inicio.date()))

    return render_template(
        "agenda/agendamento_form.html", booking=booking, laboratorios=laboratorios, turnos=TURNOS
    )


@agenda_bp.route("/<int:id>/historico")
@login_required
def historico_agendamento(id):
    booking = AgendamentoLaboratorio.query.get_or_404(id)
    if not can_manage_booking(booking):
        abort(403)
    historicos = (
        HistoricoAgendamentoLaboratorio.query.options(joinedload(HistoricoAgendamentoLaboratorio.usuario))
        .filter_by(agendamento_id=booking.id)
        .order_by(HistoricoAgendamentoLaboratorio.criado_em.desc())
        .all()
    )
    return render_template("agenda/historico.html", booking=booking, historicos=historicos)


@agenda_bp.route("/<int:id>/assumir", methods=["POST"])
@roles_required("admin", "tecnico")
def assumir_agendamento(id):
    booking = AgendamentoLaboratorio.query.with_for_update().get_or_404(id)
    try:
        if booking.status != "solicitado":
            raise ValidationError("Somente solicitações pendentes podem ser assumidas.")
        claim_booking(booking)
        db.session.commit()
    except (ValidationError, SQLAlchemyError) as error:
        db.session.rollback()
        flash(str(error) if isinstance(error, ValidationError) else "Não foi possível assumir a solicitação.", "danger")
    else:
        flash(f"{booking.protocolo} está sob seu atendimento.", "success")
    return redirect(url_for("agenda.index", data=booking.inicio.date()))


@agenda_bp.route("/<int:id>/confirmar", methods=["POST"])
@roles_required("admin", "tecnico")
def confirmar_agendamento(id):
    booking = AgendamentoLaboratorio.query.with_for_update().get_or_404(id)
    try:
        if booking.status != "solicitado":
            raise ValidationError("Somente solicitações pendentes podem ser confirmadas.")
        if booking.capacidade_reserva and booking.quantidade_alunos > booking.capacidade_reserva:
            raise ValidationError("A quantidade de alunos excede a capacidade do laboratório.")
        laboratorios_ids = [laboratorio.id for laboratorio in booking.laboratorios_reservados]
        anterior = (
            AgendamentoLaboratorio.query.filter(
                or_(
                    AgendamentoLaboratorio.laboratorio_id.in_(laboratorios_ids),
                    AgendamentoLaboratorio.laboratorio_secundario_id.in_(laboratorios_ids),
                ),
                AgendamentoLaboratorio.status == "solicitado",
                AgendamentoLaboratorio.inicio == booking.inicio,
                AgendamentoLaboratorio.fim == booking.fim,
                AgendamentoLaboratorio.id < booking.id,
            )
            .order_by(AgendamentoLaboratorio.criado_em, AgendamentoLaboratorio.id)
            .first()
        )
        if anterior:
            raise ValidationError(
                f"A solicitação {anterior.protocolo} chegou antes e deve ser analisada primeiro."
            )
        for laboratorio in booking.laboratorios_reservados:
            if booking_conflict(laboratorio.id, booking.inicio, booking.fim, booking.id):
                raise ValidationError("Este horário já foi confirmado para outra solicitação.")
        claim_booking(booking)
        booking.status = "agendado"
        booking.confirmado_por_id = current_user.id
        booking.confirmado_em = utc_now()
        booking.motivo_decisao = optional_text(request.form, "motivo_decisao")
        register_history(booking, "agendamento_confirmado", booking.motivo_decisao or "Agendamento confirmado.")
        resolve_booking_notifications(booking)
        db.session.commit()
    except (ValidationError, SQLAlchemyError) as error:
        db.session.rollback()
        flash(str(error) if isinstance(error, ValidationError) else "Não foi possível confirmar a solicitação.", "danger")
    else:
        flash(
            f"{booking.protocolo} confirmado: {booking.laboratorio_label}, {booking.inicio.strftime('%d/%m/%Y')} "
            f"no turno da {booking.turno_label.lower()}, para {booking.quantidade_alunos} alunos. "
            "Oriente o professor a chegar com antecedência e comunicar qualquer cancelamento.",
            "success",
        )
    return redirect(url_for("agenda.index", data=booking.inicio.date()))


@agenda_bp.route("/<int:id>/recusar", methods=["POST"])
@roles_required("admin", "tecnico")
def recusar_agendamento(id):
    booking = AgendamentoLaboratorio.query.get_or_404(id)
    try:
        if booking.status != "solicitado":
            raise ValidationError("Somente solicitações pendentes podem ser recusadas.")
        claim_booking(booking)
        booking.status = "recusado"
        booking.motivo_decisao = required_text(request.form, "motivo_decisao", "Motivo da recusa")
        register_history(booking, "solicitacao_recusada", booking.motivo_decisao)
        resolve_booking_notifications(booking)
        db.session.commit()
    except (ValidationError, SQLAlchemyError) as error:
        db.session.rollback()
        flash(str(error) if isinstance(error, ValidationError) else "Não foi possível recusar a solicitação.", "danger")
    else:
        flash(f"{booking.protocolo} foi recusada e ficou registrada no histórico.", "success")
    return redirect(url_for("agenda.index", data=booking.inicio.date()))


@agenda_bp.route("/<int:id>/cancelar", methods=["POST"])
@login_required
def cancelar_agendamento(id):
    booking = AgendamentoLaboratorio.query.get_or_404(id)
    if not can_manage_booking(booking):
        abort(403)
    if booking.status in STATUS_ATIVOS:
        try:
            if has_role(*SUPPORT_ROLES):
                claim_booking(booking, allow_master_override=True)
            booking.status = "cancelado"
            booking.motivo_decisao = optional_text(request.form, "motivo_decisao") or "Cancelado pelo solicitante ou pela equipe."
            register_history(booking, "agendamento_cancelado", booking.motivo_decisao)
            resolve_booking_notifications(booking)
            db.session.commit()
        except (ValidationError, SQLAlchemyError) as error:
            db.session.rollback()
            flash(str(error) if isinstance(error, ValidationError) else "Não foi possível cancelar a reserva.", "danger")
        else:
            flash("Agendamento cancelado e preservado no histórico.", "success")
    return redirect(url_for("agenda.index", data=request.form.get("data_retorno") or date.today()))


@agenda_bp.route("/<int:id>/reduzir-ilha", methods=["POST"])
@roles_required("admin", "tecnico")
def reduzir_ilha_agendamento(id):
    """Libera um dos laboratórios de uma ilha confirmada após ajuste de turma."""
    booking = AgendamentoLaboratorio.query.with_for_update().get_or_404(id)
    try:
        if booking.status != "agendado" or not booking.laboratorio_secundario_id:
            raise ValidationError("Somente uma ilha confirmada pode ser reduzida.")
        quantidade_real = integer(
            request.form, "quantidade_real", "Quantidade real de alunos", 1
        )
        laboratorios_compativeis = sorted(
            (
                laboratorio
                for laboratorio in booking.laboratorios_reservados
                if not laboratorio.capacidade or quantidade_real <= laboratorio.capacidade
            ),
            key=lambda laboratorio: (laboratorio.capacidade or 999999, laboratorio.id),
        )
        if not laboratorios_compativeis:
            raise ValidationError(
                "A quantidade real ainda exige a ilha completa (máximo de 32 alunos por laboratório)."
            )

        laboratorio_mantido = laboratorios_compativeis[0]
        laboratorio_liberado = next(
            laboratorio
            for laboratorio in booking.laboratorios_reservados
            if laboratorio.id != laboratorio_mantido.id
        )
        conflito = booking_conflict(
            laboratorio_mantido.id, booking.inicio, booking.fim, booking.id
        )
        if conflito:
            raise ValidationError("O laboratório que permanecerá já possui outra reserva confirmada.")

        ilha_anterior = booking.laboratorio_label
        quantidade_anterior = booking.quantidade_alunos
        observacao = optional_text(request.form, "observacao_reducao")
        booking.laboratorio_id = laboratorio_mantido.id
        booking.laboratorio_secundario_id = None
        booking.ilha_codigo = None
        booking.quantidade_alunos = quantidade_real
        descricao = (
            f"{ilha_anterior} reduzida para {laboratorio_mantido.nome}; "
            f"{laboratorio_liberado.nome} liberado. Turma ajustada de "
            f"{quantidade_anterior} para {quantidade_real} aluno(s)."
        )
        if observacao:
            descricao = f"{descricao} Observação: {observacao}"
        register_history(booking, "ilha_reduzida", descricao)
        db.session.commit()
    except (ValidationError, SQLAlchemyError) as error:
        db.session.rollback()
        flash(
            str(error) if isinstance(error, ValidationError) else "Não foi possível reduzir a ilha.",
            "danger",
        )
    else:
        flash(
            f"Ilha ajustada: {laboratorio_mantido.nome} permanece reservado e "
            f"{laboratorio_liberado.nome} está disponível para novo agendamento.",
            "success",
        )
    return redirect(url_for("agenda.index", data=request.form.get("data_retorno") or booking.inicio.date()))


@agenda_bp.route("/<int:id>/excluir", methods=["POST"])
@roles_required("admin", "tecnico")
def excluir_agendamento(id):
    booking = AgendamentoLaboratorio.query.get_or_404(id)
    try:
        if booking.status == "excluido":
            raise ValidationError("Este agendamento já foi excluído.")
        claim_booking(booking, allow_master_override=True)
        booking.status = "excluido"
        booking.motivo_decisao = optional_text(request.form, "motivo_decisao") or "Excluído pela equipe responsável."
        register_history(booking, "agendamento_excluido", booking.motivo_decisao)
        resolve_booking_notifications(booking)
        db.session.commit()
    except (ValidationError, SQLAlchemyError) as error:
        db.session.rollback()
        flash(str(error) if isinstance(error, ValidationError) else "Não foi possível excluir o agendamento.", "danger")
    else:
        flash("Agendamento excluído do painel ativo e preservado no histórico.", "success")
    return redirect(url_for("agenda.index", data=request.form.get("data_retorno") or date.today()))


@agenda_bp.route("/laboratorios/novo", methods=["POST"])
@roles_required("admin", "tecnico")
def novo_laboratorio():
    try:
        capacidade_raw = (request.form.get("capacidade") or "").strip()
        laboratorio = Laboratorio(
            nome=required_text(request.form, "nome", "Nome", 120),
            localizacao=optional_text(request.form, "localizacao", 150),
            capacidade=integer(request.form, "capacidade", "Capacidade", 1)
            if capacidade_raw
            else None,
            descricao=optional_text(request.form, "descricao"),
            ativo=True,
        )
        db.session.add(laboratorio)
        db.session.commit()
    except ValidationError as error:
        db.session.rollback()
        flash(str(error), "danger")
    except IntegrityError:
        db.session.rollback()
        flash("Já existe um laboratório com esse nome.", "danger")
    except SQLAlchemyError:
        db.session.rollback()
        flash("Não foi possível cadastrar o laboratório.", "danger")
    else:
        flash("Laboratório cadastrado com sucesso.", "success")
    return redirect(url_for("agenda.index", data=request.form.get("data_retorno") or date.today()))


@agenda_bp.route("/laboratorios/<int:id>/alternar", methods=["POST"])
@roles_required("admin", "tecnico")
def alternar_laboratorio(id):
    laboratorio = Laboratorio.query.get_or_404(id)
    try:
        laboratorio.ativo = not laboratorio.ativo
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        flash("Não foi possível alterar o laboratório.", "danger")
    else:
        estado = "ativado" if laboratorio.ativo else "desativado"
        flash(f"Laboratório {estado}.", "success")
    return redirect(url_for("agenda.index", data=request.form.get("data_retorno") or date.today()))


@agenda_bp.route("/recados/novo", methods=["POST"])
@roles_required("admin", "tecnico")
def novo_recado():
    try:
        data_inicio = parse_date_value(request.form.get("data_inicio"), "Data inicial")
        data_fim = parse_date_value(request.form.get("data_fim"), "Data final")
        if data_fim < data_inicio:
            raise ValidationError("A data final deve ser igual ou posterior à data inicial.")
        recado = Recado(
            titulo=required_text(request.form, "titulo", "Título", 150),
            mensagem=required_text(request.form, "mensagem", "Mensagem"),
            prioridade=choice(
                request.form, "prioridade", PRIORIDADES_RECADOS, "Prioridade"
            ),
            data_inicio=data_inicio,
            data_fim=data_fim,
            autor_id=current_user.id,
        )
        db.session.add(recado)
        db.session.commit()
    except ValidationError as error:
        db.session.rollback()
        flash(str(error), "danger")
    except SQLAlchemyError:
        db.session.rollback()
        flash("Não foi possível publicar o recado.", "danger")
    else:
        notify_general_support(
            "recado",
            "Novo recado publicado",
            f"{recado.titulo} ficará visível até {recado.data_fim.strftime('%d/%m/%Y')}.",
            url_for("agenda.recados"),
        )
        db.session.commit()
        flash("Recado publicado com sucesso.", "success")
    return redirect(url_for("agenda.recados", data=request.form.get("data_retorno") or date.today()))


@agenda_bp.route("/recados/<int:id>/excluir", methods=["POST"])
@roles_required("admin", "tecnico")
def excluir_recado(id):
    recado = Recado.query.get_or_404(id)
    try:
        db.session.delete(recado)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        flash("Não foi possível excluir o recado.", "danger")
    else:
        flash("Recado excluído.", "success")
    return redirect(url_for("agenda.recados", data=request.form.get("data_retorno") or date.today()))


@agenda_bp.route("/recados")
@roles_required("admin", "tecnico")
def recados():
    hoje = date.today()
    recados_publicados = (
        Recado.query.options(joinedload(Recado.autor))
        .order_by(Recado.data_inicio.desc(), Recado.criado_em.desc())
        .all()
    )
    return render_template("agenda/recados.html", recados=recados_publicados, hoje=hoje)
